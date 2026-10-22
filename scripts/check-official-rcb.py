"""
Fills the official Cemig RCB workbook (docs/editais/Planilha-RCB-CPP-2026-rev01.xlsx)
with an example project from prototype/seed and lets LibreOffice recalculate it,
so the app's results can be compared with the utility's own spreadsheet.

Usage: python3 scripts/check-official-rcb.py p1 [p2 ...]
Needs openpyxl and LibreOffice (soffice). Only lighting and air conditioning
actions are supported; other end uses need data the prototype does not hold.
"""
import json
import subprocess
import sys
import tempfile
import warnings
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")
ROOT = Path(__file__).resolve().parent.parent
WORKBOOK = ROOT / "docs/editais/Planilha-RCB-CPP-2026-rev01.xlsx"
SUBCLASS = {
    "A2": "A2 - De 88 kV a 230 kV",
    "A3": "A3 - 69 kV",
    "A3a": "A3a - De 30 kV a 44 kV",
    "A4": "A4 - De 2,3 kV a 25 kV",
    "AS": "AS - Subterrâneo",
    "B1": "B1 - Baixa tensão (residencial)",
    "B2": "B2 - Baixa tensão (rural)",
    "B3": "B3 - Baixa tensão (demais classes)",
    "B4": "B4 - Iluminação Pública",
}
# Columns H, I, J... hold one equipment each in the benefit sheets.
COLUMNS = [openpyxl.utils.get_column_letter(c) for c in range(8, 40)]
BTU = 12000  # nominal capacity used to express air conditioner power as BTU/COP


def fill(project):
    wb = openpyxl.load_workbook(WORKBOOK)
    costs = project.get("custos", {})
    n = lambda v: float(v or 0)
    wb["Projeto"]["J24"] = SUBCLASS[project["nivelTensao"]]
    rows = {"Iluminação": 8, "Climatização": 8}
    cols = {"Iluminação": 0, "Climatização": 0}
    for a in project["acoes"]:
        use = a["uso"]
        if use not in rows:
            raise SystemExit(f"Unsupported end use: {use}")
        cost_sheet = "IlumCusto" if use == "Iluminação" else "CondAmbCusto"
        r = rows[use]
        ws = wb[cost_sheet]
        ws[f"C{r}"], ws[f"H{r}"], ws[f"I{r}"], ws[f"J{r}"] = a.get("desc", use), n(a["vida"]), n(a["qtd"]), n(a["custo"])
        rows[use] += 1
        col = COLUMNS[cols[use]]
        cols[use] += 1
        hours_per_day = n(a["horas"]) / 365
        if use == "Iluminação":
            b = wb["IlumBenef"]
            current = {9: n(a["pa"]), 10: n(a["qtd"]), 11: 0, 12: 0, 14: hours_per_day, 15: 365, 17: n(a["nup"]), 18: n(a["nd"]), 19: n(a["nm"])}
            proposed = {30: n(a["pd"]), 31: n(a["qtd"]), 32: 0, 33: 0, 35: hours_per_day, 36: 365, 38: n(a["nup"]), 39: n(a["nd"]), 40: n(a["nm"])}
        else:
            b = wb["CondAmbBenef"]
            current = {8: BTU, 9: BTU * 0.293 / n(a["pa"]), 10: n(a["qtd"]), 12: 1, 14: hours_per_day, 15: 365, 17: n(a["nup"]), 18: n(a["nd"]), 19: n(a["nm"])}
            proposed = {28: BTU, 29: BTU * 0.293 / n(a["pd"]), 30: n(a["qtd"]), 32: 1, 34: hours_per_day, 35: 365, 37: n(a["nup"]), 38: n(a["nd"]), 39: n(a["nm"])}
        for row, value in {**current, **proposed}.items():
            b[f"{col}{row}"] = value
    il = wb["IlumCusto"]
    il["C116"], il["I116"], il["J116"] = "Projeto executivo e instalação", 1, n(costs.get("projExec")) + n(costs.get("motOutros"))
    il["C142"], il["I142"], il["J142"] = "Descarte", 1, n(costs.get("desc"))
    t = wb["Terceiros"]
    t["J7"], t["J8"], t["J9"] = n(costs.get("diag")), n(costs.get("gerenc")), n(costs.get("relFinal"))
    wb["Marketing"]["I7"], wb["Marketing"]["J7"] = 1, n(costs.get("mkt"))
    wb["Treinamento"]["I7"], wb["Treinamento"]["J7"] = 1, n(costs.get("trein"))
    wb["M&V"]["I9"] = n(costs.get("mv"))
    return wb


def read(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    r, cc = wb["RCB"], wb["CustoContábil"]
    uses = {}
    for row, name in [(9, "Iluminação"), (10, "Climatização")]:
        cat, bat = r[f"E{row}"].value or 0, r[f"F{row}"].value or 0
        if bat:
            uses[name] = {"cat": cat, "bat": bat, "rcb": cat / bat, "ee": r[f"C{row}"].value, "rdp": r[f"D{row}"].value}
    cat = sum(u["cat"] for u in uses.values())
    bat = sum(u["bat"] for u in uses.values())
    return {"pee": cc["F18"].value, "audit": cc["F10"].value, "rcb": cat / bat, "uses": uses}


def main(ids):
    out = {}
    with tempfile.TemporaryDirectory() as tmp:
        for pid in ids:
            project = json.loads((ROOT / f"prototype/seed/projetos/{pid}.json").read_text())
            src = Path(tmp) / f"{pid}.xlsx"
            fill(project).save(src)
            subprocess.run(
                ["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(Path(tmp) / "out"), str(src)],
                check=True, capture_output=True, timeout=600,
            )
            out[pid] = read(Path(tmp) / "out" / f"{pid}.xlsx")
    print(json.dumps(out, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main(sys.argv[1:] or ["p1"])
